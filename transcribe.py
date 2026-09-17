"""Groq Whisper transcription for recorded files, with timecoded .docx output."""
import json
import re
import subprocess
import tempfile
from pathlib import Path

import requests
from docx import Document

import config
import ffmpeg_utils

GROQ_URL = "https://api.groq.com/openai/v1/audio/transcriptions"
MAX_UPLOAD_BYTES = 24 * 1024 * 1024
CREATE_NO_WINDOW = 0x08000000
# Groq answers 403 to every request (even unauthenticated) from blocked regions, incl. Russia.
GEO_BLOCK_MESSAGE = "Groq недоступен из вашего региона (ошибка 403) — включите VPN и повторите"

SPEAKER_ME = "Я"
SPEAKER_OTHER = "Собеседник"

# Whisper invents these on silence (subtitle credits from its training data).
HALLUCINATION_RE = re.compile(
    r"продолжение следует|субтитр|спасибо за просмотр|подписывайтесь на канал|редактор|корректор",
    re.IGNORECASE,
)


class TranscriptionError(Exception):
    pass


def _audio_stream_titles(ffprobe_path: str, source: Path) -> list[str]:
    result = subprocess.run(
        [ffprobe_path, "-v", "error", "-select_streams", "a",
         "-show_entries", "stream=index:stream_tags", "-of", "json", str(source)],
        capture_output=True,
        stdin=subprocess.DEVNULL,
        creationflags=CREATE_NO_WINDOW,
    )
    try:
        streams = json.loads(result.stdout.decode("utf-8", errors="replace")).get("streams", [])
    except json.JSONDecodeError:
        return []
    # ffmpeg writes a stream "title" into MP4 but reads it back as "name".
    return [s.get("tags", {}).get("title") or s.get("tags", {}).get("name") or "" for s in streams]


def _get_duration_seconds(ffprobe_path: str, source: Path) -> float:
    result = subprocess.run(
        [ffprobe_path, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(source)],
        capture_output=True,
        stdin=subprocess.DEVNULL,
        creationflags=CREATE_NO_WINDOW,
    )
    try:
        return float(result.stdout.strip())
    except ValueError:
        return 0.0


def _extract_compact_audio(ffmpeg_path: str, source: Path, dest: Path, stream_index: int) -> None:
    cmd = [
        ffmpeg_path, "-y", "-hide_banner", "-loglevel", "warning",
        "-i", str(source),
        "-map", f"0:a:{stream_index}",
        "-ac", "1", "-ar", "16000", "-c:a", "aac", "-b:a", "64k",
        str(dest),
    ]
    result = subprocess.run(cmd, capture_output=True, stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW)
    if result.returncode != 0:
        raise TranscriptionError(f"ffmpeg audio extraction failed: {result.stderr[-1500:].decode(errors='replace')}")


def _split_audio(ffmpeg_path: str, source: Path, out_dir: Path, segment_seconds: int = 600) -> list[Path]:
    pattern = out_dir / f"{source.stem}_part_%03d.m4a"
    cmd = [
        ffmpeg_path, "-y", "-hide_banner", "-loglevel", "warning",
        "-i", str(source),
        "-f", "segment", "-segment_time", str(segment_seconds),
        "-c", "copy",
        str(pattern),
    ]
    result = subprocess.run(cmd, capture_output=True, stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW)
    if result.returncode != 0:
        raise TranscriptionError(f"ffmpeg split failed: {result.stderr[-1500:].decode(errors='replace')}")
    return sorted(out_dir.glob(f"{source.stem}_part_*.m4a"))


def _post_to_groq(file_path: Path, api_key: str):
    with open(file_path, "rb") as f:
        return requests.post(
            GROQ_URL,
            headers={"Authorization": f"Bearer {api_key}"},
            files={"file": (file_path.name, f, "audio/mp4")},
            data={
                "model": config.GROQ_MODEL,
                "response_format": "verbose_json",
                "temperature": "0",
                "timestamp_granularities[]": ["segment", "word"],
            },
            timeout=300,
        )


def _is_hallucination(segment: dict) -> bool:
    if HALLUCINATION_RE.search(segment.get("text", "")):
        return True
    return segment.get("no_speech_prob", 0.0) >= 0.6 and segment.get("avg_logprob", 0.0) <= -0.5


def _call_groq(file_path: Path) -> list[dict]:
    """Returns a list of {"start": float, "end": float, "text": str} segments."""
    response = _post_to_groq(file_path, config.GROQ_API_KEY)
    if response.status_code == 429 and config.GROQ_API_KEY_2:
        response = _post_to_groq(file_path, config.GROQ_API_KEY_2)
    if response.status_code == 403:
        raise TranscriptionError(GEO_BLOCK_MESSAGE)
    if response.status_code != 200:
        raise TranscriptionError(f"Groq API error {response.status_code}: {response.text[:500]}")

    data = response.json()
    segments = data.get("segments")
    if segments is None:
        text = (data.get("text") or "").strip()
        return [{"start": 0.0, "end": 0.0, "text": text}] if text else []

    # Whisper stretches a segment's start back over preceding silence; the first
    # word's timestamp is where speech actually begins. This matters when two
    # tracks are merged into one dialogue by time.
    word_starts = sorted(float(w.get("start", 0.0)) for w in data.get("words") or [])

    result = []
    for s in segments:
        text = s.get("text", "").strip()
        if not text or _is_hallucination(s):
            continue
        start, end = float(s.get("start", 0.0)), float(s.get("end", 0.0))
        first_word = next((w for w in word_starts if start <= w < end), None)
        result.append({"start": first_word if first_word is not None else start, "end": end, "text": text})
    return result


def _transcribe_track(ffmpeg_path: str, ffprobe_path: str, source: Path, stream_index: int, tmp_dir: Path) -> list[dict]:
    compact_audio = tmp_dir / f"track{stream_index}.m4a"
    _extract_compact_audio(ffmpeg_path, source, compact_audio, stream_index)

    if compact_audio.stat().st_size <= MAX_UPLOAD_BYTES:
        parts = [compact_audio]
    else:
        parts = _split_audio(ffmpeg_path, compact_audio, tmp_dir)

    segments = []
    offset = 0.0
    for part in parts:
        for segment in _call_groq(part):
            segments.append({
                "start": segment["start"] + offset,
                "end": segment["end"] + offset,
                "text": segment["text"],
            })
        offset += _get_duration_seconds(ffprobe_path, part)
    return segments


def _format_timecode(seconds: float) -> str:
    total = int(seconds)
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _segment_line(segment: dict) -> str:
    speaker = f"{segment['speaker']}: " if segment.get("speaker") else ""
    return f"[{_format_timecode(segment['start'])}] {speaker}{segment['text']}"


def segments_to_text(segments: list[dict]) -> str:
    return "\n".join(_segment_line(s) for s in segments)


def has_speakers(segments: list[dict]) -> bool:
    return any(s.get("speaker") for s in segments)


def _add_markdown(doc: Document, markdown: str) -> None:
    for raw_line in markdown.splitlines():
        line = raw_line.strip().replace("**", "")
        if not line:
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=3)
        elif line[:2] in ("- ", "* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line[0].isdigit() and ". " in line[:4]:
            doc.add_paragraph(line.split(". ", 1)[1].strip(), style="List Number")
        else:
            doc.add_paragraph(line)


def build_docx(
    recording_path: Path,
    segments: list[dict],
    summary_markdown: str | None = None,
    summary_note: str | None = None,
) -> Path:
    doc = Document()
    doc.add_heading(recording_path.stem, level=1)

    doc.add_heading("Саммари встречи", level=2)
    if summary_markdown:
        _add_markdown(doc, summary_markdown)
    elif summary_note:
        doc.add_paragraph(summary_note)

    doc.add_heading("Полная расшифровка", level=2)
    if has_speakers(segments):
        doc.add_paragraph(
            f"«{SPEAKER_ME}» — автор записи (микрофон), «{SPEAKER_OTHER}» — все остальные участники звонка."
        )
    for segment in segments:
        doc.add_paragraph(_segment_line(segment))

    docx_path = config.TRANSCRIPTS_DIR / f"{recording_path.stem}_transcript.docx"
    try:
        doc.save(docx_path)
    except PermissionError:
        # The first version is probably open in Word — don't lose the summary.
        docx_path = docx_path.with_name(f"{recording_path.stem}_transcript_summary.docx")
        doc.save(docx_path)
    return docx_path


def transcribe_segments(recording_path: Path) -> list[dict]:
    if not config.GROQ_API_KEY:
        raise TranscriptionError("GROQ_API_KEY is not configured")

    ffmpeg_path = ffmpeg_utils.find_ffmpeg()
    ffprobe_path = ffmpeg_utils.find_ffprobe(ffmpeg_path)

    titles = _audio_stream_titles(ffprobe_path, recording_path)
    if not titles:
        raise TranscriptionError("Recording has no audio track (no microphone) — nothing to transcribe")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        if config.MIC_TRACK_TITLE in titles and config.SYSTEM_TRACK_TITLE in titles:
            segments = []
            for title, speaker in ((config.MIC_TRACK_TITLE, SPEAKER_ME), (config.SYSTEM_TRACK_TITLE, SPEAKER_OTHER)):
                for segment in _transcribe_track(ffmpeg_path, ffprobe_path, recording_path, titles.index(title), tmp_dir):
                    segments.append({**segment, "speaker": speaker})
            segments.sort(key=lambda s: s["start"])
        else:
            segments = _transcribe_track(ffmpeg_path, ffprobe_path, recording_path, 0, tmp_dir)

    if not segments:
        raise TranscriptionError("Groq returned an empty transcript")

    return segments
